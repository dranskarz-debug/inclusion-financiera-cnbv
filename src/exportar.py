"""Exporta el Excel limpio con diccionario de datos y bitácora."""

from __future__ import annotations

import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SECTOR = {"total": "Todos los sectores", "bm": "Banca múltiple", "bd": "Banca de desarrollo",
          "socap": "Sociedades Cooperativas de Ahorro y Préstamo", "sofipo": "Sociedades Financieras Populares"}
GRUPO = {"banca": "Banca (múltiple + desarrollo)", "eacp": "EACP (SOCAP + SOFIPO)",
         "infraestructura": "Infraestructura", "transacciones": "Transacciones"}
FIJAS = {
    "clave_mun": "Clave INEGI del municipio (estado × 1000 + municipio). 99999 = no identificado",
    "clave_edo": "Clave INEGI del estado",
    "region": "Región según la CNBV",
    "estado": "Entidad federativa",
    "municipio": "Nombre del municipio",
    "poblacion_total": "Población total",
    "poblacion_adulta": "Población de 15 años o más",
    "poblacion_adulta_mujeres": "Población adulta, mujeres",
    "poblacion_adulta_hombres": "Población adulta, hombres",
    "tipo_de_poblacion": "Clasificación del municipio por tamaño de población (CNBV)",
    "indice_de_rezago_social": "Índice de rezago social (CONEVAL)",
    "es_no_identificado": "Bandera: registro 99999 sin municipio asignado",
    "sin_poblacion": "Bandera: municipio sin dato de población (de creación reciente)",
    "usar_en_tasas": "Bandera: se usa en indicadores por habitante",
    "sin_punto_de_acceso": "Bandera: sin sucursal, cajero ni corresponsal",
    "sucursales_por_10k_adultos": "Sucursales por cada 10,000 adultos (calculado)",
    "cajeros_por_10k_adultos": "Cajeros automáticos por cada 10,000 adultos (calculado)",
    "corresponsales_por_10k_adultos": "Corresponsales por cada 10,000 adultos (calculado)",
    "cuentas_por_adulto": "Cuentas de captación por adulto (calculado; ver nota sobre domicilio)",
}


ACENTOS = {"automaticos": "automáticos", "unicos": "únicos", "depositos": "depósitos",
           "debito": "débito", "efectuan": "efectúan", "tpv": "TPV"}


def describir(col: str) -> tuple[str, str]:
    if col in FIJAS:
        return "Información general / calculada", FIJAS[col]
    partes = col.split("__")
    tema = "Captación" if partes[0] == "capt" else "Infraestructura y transacciones"
    if partes[0] == "capt":
        partes = partes[1:]
    sector = SECTOR.get(partes[0], partes[0])
    resto = partes[1:]
    grupo = GRUPO.get(resto[0]) if len(resto) > 1 and resto[0] in GRUPO else None
    palabras = [ACENTOS.get(p, p) for p in resto[-1].split("_")]
    variable = " ".join(palabras)
    variable = variable[0].upper() + variable[1:]
    if variable == "Cuentas total":
        variable = "Total de cuentas de captación (sin tarjetas de débito)"
    texto = f"{variable} — {sector}" + (f", {grupo}" if grupo else "")
    return tema, texto


def _formatear(ws, anchos=None):
    encabezado = PatternFill("solid", start_color="1C5CAB")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(j)].width = (anchos or {}).get(j, min(max(largo + 2, 10), 48))
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
            elif isinstance(c.value, int) and not isinstance(c.value, bool):
                c.number_format = "#,##0" if "clave" not in str(ws.cell(1, c.column).value) else "0"


def exportar(bdif, por_estado, por_rezago, por_tipo, bit_infra, bit_capt, ruta):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    dicc = pd.DataFrame([(c, *describir(c), str(bdif[c].dtype)) for c in bdif.columns],
                        columns=["columna", "tema", "descripción", "tipo"])
    notas = pd.DataFrame({"Notas": [
        "Base de Datos de Inclusión Financiera de la CNBV, corte a marzo de 2026 (datos públicos).",
        "Fuente: https://www.gob.mx/cnbv/acciones-y-programas/bases-de-datos-de-inclusion-financiera",
        "Hojas originales usadas: 'BD Infraestructura Mun' y 'BD Captación Mun', unidas por clave municipal.",
        "Los indicadores por grupo suman numeradores y denominadores antes de dividir.",
        "Se excluyen de los indicadores por habitante el registro 99999 y 3 municipios sin dato de población.",
        "Las cuentas se asignan al municipio de registro, no necesariamente al de residencia.",
        "Limpieza y análisis: Leonardo Chapa Gallardo.",
    ]})
    bitacora = pd.DataFrame(
        [("Infraestructura", l) for l in bit_infra] + [("Captación", l) for l in bit_capt],
        columns=["hoja", "validación superada"])

    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        notas.to_excel(w, sheet_name="Notas", index=False)
        bdif.to_excel(w, sheet_name="Municipios", index=False)
        por_estado.reset_index().to_excel(w, sheet_name="Por estado", index=False)
        por_rezago.reset_index().to_excel(w, sheet_name="Por rezago social", index=False)
        por_tipo.reset_index().to_excel(w, sheet_name="Por tipo de población", index=False)
        dicc.to_excel(w, sheet_name="Diccionario", index=False)
        bitacora.to_excel(w, sheet_name="Bitácora de validación", index=False)
        for nombre, ws in w.sheets.items():
            _formatear(ws, {1: 110} if nombre == "Notas" else None)
